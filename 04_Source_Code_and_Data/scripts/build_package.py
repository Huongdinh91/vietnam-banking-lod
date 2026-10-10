"""Build the final dataset from the group's submission TTL.

Input : scripts/input_vietnam-banking-submission.ttl
Output: data/vietnam-banking.ttl, data/csv/*.csv, data/vietnam_banking_lod_dataset.xlsx
Run   : python scripts/build_package.py   (from the package root)
"""
import csv, os
from rdflib import Graph, Namespace, URIRef, Literal
from rdflib.namespace import RDF, RDFS, OWL, XSD, DCTERMS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IN = os.path.join(ROOT, "scripts", "input_vietnam-banking-submission.ttl")
OUT_TTL = os.path.join(ROOT, "data", "vietnam-banking.ttl")
CSV_DIR = os.path.join(ROOT, "data", "csv")
BUILD_DATE = "2026-10-09"

B = "https://data.example.org/vietnam-banking/"
E = Namespace(B)
BANK, LOC, BT, SV = (Namespace(B + x + "/") for x in ("bank", "location", "bank-type", "service"))
VOID = Namespace("http://rdfs.org/ns/void#")
SCHEMA = Namespace("https://schema.org/")
WD = "http://www.wikidata.org/entity/"
DBR = "http://dbpedia.org/resource/"

g = Graph(); g.parse(IN, format="turtle")
for p, ns in [("", E), ("void", VOID), ("schema", SCHEMA), ("dcterms", DCTERMS), ("owl", OWL)]:
    g.bind(p, ns, override=True)

def drop(node):
    for t in list(g.triples((node, None, None))) + list(g.triples((None, None, node))):
        g.remove(t)

# 1. Orphan locations (no bank uses them) and their external IRIs
for n in [LOC["vinh"], LOC["da-nang"], URIRef(DBR + "Vinh"), URIRef(DBR + "Da_Nang"),
          URIRef(WD + "Q194292"), URIRef(WD + "Q25282")]:
    drop(n)
# 2. Leftovers of links removed earlier (DBpedia Sacombank = 404, VIB = unverified)
for n in [URIRef(DBR + "Sacombank"), URIRef(DBR + "Vietnam_International_Commercial_Joint_Stock_Bank")]:
    drop(n)
# 3. GPBank had no bank type (violates qualifiedCardinality 1)
g.add((BANK["gpbank"], E.hasBankType, BT["private-domestic-commercial"]))
g.add((BT["private-domestic-commercial"], RDFS.comment, Literal(
    "Ngân hàng thương mại TNHH MTV trong nước do một NHTM cổ phần sở hữu sau chuyển giao bắt buộc (MBV thuộc MB, GPBank thuộc VPBank).", lang="vi")))

# 4. Additional links found by searching wikidata.org (label match) and one DBpedia page check
NEW_LINKS = {  # bank slug: [(external IRI, method, evidence, confidence)]
 "vpbank":    [(WD+"Q20025914","wikidata-search","https://www.wikidata.org/wiki/Q20025914","medium")],
 "mbbank":    [(WD+"Q10800269","wikidata-search","https://www.wikidata.org/wiki/Q10800269","medium")],
 "hdbank":    [(WD+"Q10800273","wikidata-search","https://www.wikidata.org/wiki/Q10800273","medium")],
 "lpbank":    [(WD+"Q15729124","wikidata-search","https://www.wikidata.org/wiki/Q15729124","medium")],
 "seabank":   [(WD+"Q10800276","wikidata-search","https://www.wikidata.org/wiki/Q10800276","medium")],
 "msb":       [(WD+"Q10800277","wikidata-search","https://www.wikidata.org/wiki/Q10800277","medium")],
 "namabank":  [(WD+"Q10800266","wikidata-search","https://www.wikidata.org/wiki/Q10800266","medium")],
 "abbank":    [(WD+"Q10800288","wikidata-search","https://www.wikidata.org/wiki/Q10800288","medium")],
 "pvcombank": [(WD+"Q65171780","wikidata-search","https://www.wikidata.org/wiki/Q65171780","medium")],
 "tpbank":    [(WD+"Q7801033","wikidata-search","https://www.wikidata.org/wiki/Q7801033","medium")],
 "vib":       [(WD+"Q49054349","wikidata-search","https://www.wikidata.org/wiki/Q49054349","medium")],
 "scb":       [(WD+"Q101575325","wikidata-search","https://www.wikidata.org/wiki/Q101575325","medium")],
 "vbsp":      [(WD+"Q10800262","wikidata-search","https://www.wikidata.org/wiki/Q10800262","medium")],
 "vdb":       [(WD+"Q10800267","wikidata-search","https://www.wikidata.org/wiki/Q10800267","medium")],
 "eximbank":  [(WD+"Q10800271","wikidata-search","https://www.wikidata.org/wiki/Q10800271","low"),
               (DBR+"Eximbank_(Vietnam)","dbpedia-page-check","https://dbpedia.org/page/Eximbank_(Vietnam)","high")],
}
for slug, links in NEW_LINKS.items():
    for iri, *_ in links:
        g.add((BANK[slug], OWL.sameAs, URIRef(iri)))
        g.add((URIRef(iri), RDF.type, OWL.NamedIndividual))

# Provenance of links that were already in the submission
OLD_PROV = {
 (DBR+"Vietcombank"):("dbpedia-page-check","https://dbpedia.org/page/Vietcombank","high"),
 (DBR+"Bank_for_Investment_and_Development_of_Vietnam"):("dbpedia-page-check","https://dbpedia.org/page/BIDV","high"),
 (DBR+"Vietinbank"):("dbpedia-page-check","https://dbpedia.org/page/VietinBank","high"),
 (DBR+"Agribank_(Vietnam)"):("dbpedia-page-check","https://dbpedia.org/page/Vietnam_Bank_for_Agriculture_and_Rural_Development","high"),
 (DBR+"Orient_Commercial_Joint_Stock_Bank"):("dbpedia-page-check","https://dbpedia.org/page/Orient_Commercial_Bank","high"),
 (DBR+"Asia_Commercial_Bank"):("wikipedia-title","https://en.wikipedia.org/wiki/Asia_Commercial_Bank","medium"),
 (DBR+"Techcombank"):("wikipedia-title","https://en.wikipedia.org/wiki/Techcombank","medium"),
 (WD+"Q6122995"):("dbpedia-sameAs","https://dbpedia.org/page/Vietcombank","high"),
 (WD+"Q1003180"):("dbpedia-sameAs","https://dbpedia.org/page/BIDV","high"),
 (WD+"Q7928459"):("dbpedia-sameAs","https://dbpedia.org/page/VietinBank","high"),
 (WD+"Q1924723"):("dbpedia-sameAs","https://dbpedia.org/page/Vietnam_Bank_for_Agriculture_and_Rural_Development","high"),
 (WD+"Q7102288"):("dbpedia-sameAs","https://dbpedia.org/page/Orient_Commercial_Bank","high"),
 (WD+"Q727019"):("wikidata-search","https://www.wikidata.org/wiki/Q727019","medium"),
 (WD+"Q10541776"):("wikidata-search","https://www.wikidata.org/wiki/Q10541776","medium"),
 (WD+"Q124465825"):("wikidata-search","https://www.wikidata.org/wiki/Q124465825","medium"),
 (WD+"Q6123772"):("wikidata-search","https://www.wikidata.org/wiki/Q6123772","medium"),
 (WD+"Q881"):("wikidata-search","https://www.wikidata.org/wiki/Q881","high"),
 (WD+"Q1858"):("wikidata-search","https://www.wikidata.org/wiki/Q1858","high"),
 (WD+"Q1854"):("wikidata-search","https://www.wikidata.org/wiki/Q1854","high"),
 (DBR+"Hanoi"):("dbpedia-page-check","https://dbpedia.org/page/Hanoi","high"),
 (DBR+"Ho_Chi_Minh_City"):("dbpedia-page-check","https://dbpedia.org/page/Ho_Chi_Minh_City","high"),
 (DBR+"Vietnam"):("wikipedia-title","https://en.wikipedia.org/wiki/Vietnam","high"),
}
PROV = dict(OLD_PROV)
for links in NEW_LINKS.values():
    for iri, m, ev, c in links: PROV[iri] = (m, ev, c)

# 5. VoID / dataset metadata
D = E["dataset"]
for p in (VOID.triples, DCTERMS.modified):
    g.remove((D, p, None))
def internal_links(prefix):
    return [(s, o) for s, _, o in g.triples((None, OWL.sameAs, None)) if str(s).startswith(B) and str(o).startswith(prefix)]
for ls, prefix in [(E["linkset/wikidata"], WD), (E["linkset/dbpedia"], DBR)]:
    g.remove((ls, VOID.triples, None))
    g.add((ls, VOID.triples, Literal(len(internal_links(prefix)))))
meta = [(DCTERMS.modified, Literal(BUILD_DATE, datatype=XSD.date)),
        (DCTERMS.created, Literal("2026-10-06", datatype=XSD.date)),
        (DCTERMS.license, URIRef("https://creativecommons.org/licenses/by-sa/4.0/")),
        (DCTERMS.source, URIRef("https://en.wikipedia.org/wiki/List_of_banks_in_Vietnam")),
        (DCTERMS.creator, Literal("Nhóm 24, IT6390E Semantic Web, Đại học Bách khoa Hà Nội", lang="vi")),
        (VOID.uriSpace, Literal(B)),
        (VOID.vocabulary, URIRef(B)), (VOID.vocabulary, URIRef("https://schema.org/")),
        (VOID.vocabulary, URIRef("http://purl.org/dc/terms/")),
        (VOID.dataDump, URIRef(B + "vietnam-banking.ttl")),
        (VOID.exampleResource, BANK["vietcombank"])]
for p, o in meta:
    g.add((D, p, o))
g.remove((URIRef(B), DCTERMS.modified, None))
g.add((URIRef(B), DCTERMS.modified, Literal(BUILD_DATE, datatype=XSD.date)))
g.add((D, VOID.triples, Literal(0)))
n = len(g); g.remove((D, VOID.triples, Literal(0))); g.add((D, VOID.triples, Literal(n)))
assert len(g) == n
g.serialize(OUT_TTL, format="turtle")

# 6. Export CSV tables from the same graph (single source of truth)
def lit(s, p, lang=None):
    for o in g.objects(s, p):
        if lang is None or getattr(o, "language", None) == lang:
            return str(o)
    return ""
def slug(u): return str(u).rstrip("/").split("/")[-1]
def write(name, header, rows):
    with open(os.path.join(CSV_DIR, name), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)

banks = sorted(set(g.subjects(RDF.type, E.Bank)), key=slug)
banks = [b for b in banks if str(b).startswith(str(BANK))]
write("banks.csv", ["bank_id","internal_uri","name_vi","name_en","abbreviation","bank_type_id","headquarters_location_id",
      "country_location_id","founded_year","official_website","source_url","retrieved_date","license"],
      [[lit(b,E.bankId), str(b), lit(b,E.nameVi), lit(b,E.nameEn), lit(b,E.abbreviation),
        slug(g.value(b,E.hasBankType)), slug(g.value(b,E.hasHeadquarters)), slug(g.value(b,E.operatesInCountry)),
        lit(b,E.foundedYear), lit(b,E.officialWebsite), str(g.value(b,DCTERMS.source) or ""), lit(b,E.retrievedDate), "CC-BY-SA-4.0"]
       for b in banks])
def labelled(cls, ns):
    return sorted([s for s in g.subjects(RDF.type, cls) if str(s).startswith(str(ns))], key=slug)
write("bank_types.csv", ["bank_type_id","internal_uri","label_vi","label_en","comment_vi","bank_count"],
      [[slug(t), str(t), lit(t,RDFS.label,"vi"), lit(t,RDFS.label,"en"), lit(t,RDFS.comment,"vi"),
        len(set(g.subjects(E.hasBankType,t)))] for t in labelled(E.BankType, BT)])
locs = labelled(E.Location, LOC)
def ext(s, prefix): return next((str(o) for o in g.objects(s, OWL.sameAs) if str(o).startswith(prefix)), "")
write("locations.csv", ["location_id","internal_uri","name_vi","name_en","wikidata_uri","dbpedia_uri","banks_headquartered"],
      [[slug(l), str(l), lit(l,RDFS.label,"vi"), lit(l,RDFS.label,"en"), ext(l,WD), ext(l,DBR),
        len(set(g.subjects(E.hasHeadquarters,l)))] for l in locs])
svcs = labelled(E.Service, SV)
write("services.csv", ["service_id","internal_uri","label_vi","label_en"],
      [[slug(s), str(s), lit(s,RDFS.label,"vi"), lit(s,RDFS.label,"en")] for s in svcs])
write("bank_services.csv", ["bank_id","service_id","mapping_basis"],
      [[slug(b), slug(s), "category_based (theo loại ngân hàng, chưa đối chiếu từng website)"]
       for b in banks for s in sorted(g.objects(b,E.providesService), key=slug)])
rows = []
for s, _, o in sorted(g.triples((None, OWL.sameAs, None))):
    if not str(s).startswith(B): continue
    m, ev, c = PROV.get(str(o), ("", "", ""))
    rows.append([slug(s), "bank" if "/bank/" in str(s) else "location", "owl:sameAs",
                 "Wikidata" if str(o).startswith(WD) else "DBpedia", str(o), m, ev, c,
                 "2026-10-06" if str(o) in OLD_PROV else BUILD_DATE])
write("external_links.csv", ["subject_id","subject_type","predicate","target_dataset","external_uri","method","evidence_url","confidence","verified_date"], rows)
write("sources.csv", ["source_id","title","url","used_for","license","retrieved_date"], [
 ["SRC_WIKI_BANK_LIST","List of banks in Vietnam","https://en.wikipedia.org/wiki/List_of_banks_in_Vietnam","Tên, loại, trụ sở, website","CC BY-SA 4.0","2026-10-06"],
 ["SRC_WIKIDATA","Wikidata","https://www.wikidata.org/","IRI thực thể (owl:sameAs)","CC0 1.0","2026-10-09"],
 ["SRC_DBPEDIA","DBpedia","https://www.dbpedia.org/","IRI thực thể (owl:sameAs)","CC BY-SA","2026-10-06"],
 ["SRC_OFFICIAL_WEBSITES","Website chính thức của ngân hàng","(cột official_website)","Website; ánh xạ dịch vụ khởi đầu","Chỉ trích dẫn dữ kiện","2026-10-06"]])
write("licenses.csv", ["license_id","name","url","applies_to"], [
 ["CC-BY-SA-4.0","Creative Commons Attribution-ShareAlike 4.0","https://creativecommons.org/licenses/by-sa/4.0/","Toàn bộ dataset (kế thừa ShareAlike từ Wikipedia/DBpedia)"],
 ["CC0-1.0","Creative Commons Zero 1.0","https://creativecommons.org/publicdomain/zero/1.0/","IRI và dữ kiện lấy từ Wikidata"]])
bl = {s for s, _ in internal_links(WD) + internal_links(DBR) if "/bank/" in str(s)}
write("manifest.csv", ["metric","value"], [
 ["dataset_title","Vietnam Banking Linked Open Data"],["version","1.2.0"],["base_uri",B],
 ["created_date","2026-10-06"],["modified_date",BUILD_DATE],["triples",n],
 ["bank_count",len(banks)],["bank_type_count",len(labelled(E.BankType,BT))],["location_count",len(locs)],
 ["service_type_count",len(svcs)],["bank_service_links",sum(1 for _ in g.triples((None,E.providesService,None)))],
 ["banks_with_external_links",len(bl)],["bank_links_wikidata",sum(1 for s,_ in internal_links(WD) if "/bank/" in str(s))],
 ["bank_links_dbpedia",sum(1 for s,_ in internal_links(DBR) if "/bank/" in str(s))],
 ["location_link_coverage",f"{sum(1 for l in locs if ext(l,WD) and ext(l,DBR))}/{len(locs)}"],
 ["license","CC-BY-SA-4.0"],["encoding","UTF-8 with BOM"]])

# 7. Excel workbook with every table
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); wb.remove(wb.active)
for name in ["manifest","banks","bank_types","locations","services","bank_services","external_links","sources","licenses"]:
    ws = wb.create_sheet(name)
    with open(os.path.join(CSV_DIR, name + ".csv"), encoding="utf-8-sig") as f:
        for i, row in enumerate(csv.reader(f)):
            ws.append([int(x) if x.isdigit() and len(x) < 10 else x for x in row])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(name="Arial")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"
wb.save(os.path.join(ROOT, "data", "vietnam_banking_lod_dataset.xlsx"))
print("triples", n, "| banks", len(banks), "| linked banks", len(bl), "| locations", len(locs))
